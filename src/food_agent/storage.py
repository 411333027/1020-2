"""存照片、寫 Excel、算當日累計。"""

from __future__ import annotations

import io
import logging
import os
import time
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from PIL import Image

from . import guards
from .analyzer import Analysis
from .config import EXCEL_MAX_RETRIES, EXCEL_RETRY_SECONDS, Config

log = logging.getLogger("food_agent.storage")

HEADERS = [
    "日期時間",
    "上傳者",
    "食物敘述",
    "成份",
    "重量(公克)",
    "熱量(大卡)",
    "當日已攝取熱量(大卡)",
    "照片存放位置",
]
NOT_FOOD_LABEL = "不是食物"
STAMP_FORMAT = "%Y-%m-%d %H-%M-%S"


class StorageError(Exception):
    """存照片或寫 Excel 失敗。"""


@dataclass(frozen=True)
class Meal:
    """Excel 裡的一列，查詢時使用。"""

    when: datetime
    uploader: str
    description: str
    kcal: int | None


# ---------------------------------------------------------------------------
# 照片
# ---------------------------------------------------------------------------
def validate_image(data: bytes) -> None:
    """確認內容真的是圖片。不是的話丟 StorageError。"""
    try:
        with Image.open(io.BytesIO(data)) as image:
            image.verify()
    except Exception as exc:  # Pillow 對壞檔會丟各式各樣的例外
        raise StorageError(f"不是有效的圖片：{exc}") from exc


def save_photo(config: Config, data: bytes, when: datetime, name: str) -> Path:
    """把照片轉成 JPEG 存進 photos 資料夾，回傳實際存下的路徑。

    檔名：「YYYY-MM-DD hh-mm-ss 食物名稱簡稱.jpg」，同一秒撞名時在後面加 -2、-3。
    """
    short_name = " ".join(name.split())[: guards.MAX_NAME_LENGTH] or "food"
    stamp = when.strftime(STAMP_FORMAT)

    try:
        with Image.open(io.BytesIO(data)) as image:
            rgb = image.convert("RGB")
    except Exception as exc:
        raise StorageError(f"無法讀取圖片：{exc}") from exc

    for attempt in range(1, 100):
        try:
            path = guards.build_photo_path(config.photos_dir, stamp, short_name, attempt)
        except ValueError as exc:
            raise StorageError(str(exc)) from exc
        try:
            # "xb"：只允許建立新檔，絕不覆蓋已存在的照片
            with open(path, "xb") as handle:
                rgb.save(handle, "JPEG", quality=90)
            return path
        except FileExistsError:
            continue
        except OSError as exc:
            raise StorageError(f"存照片失敗：{exc}") from exc
    raise StorageError("同一秒撞名太多次，放棄存檔")


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def _with_retry(action, what: str):
    """Excel 被你打開時 Windows 會鎖住檔案：印警告並重試，最後仍失敗才丟 StorageError。"""
    for attempt in range(1, EXCEL_MAX_RETRIES + 1):
        try:
            return action()
        except PermissionError:
            log.warning(
                "⚠️ %s失敗：food_log.xlsx 可能正被 Excel 開著，請先關閉它。%d 秒後重試（第 %d/%d 次）",
                what, EXCEL_RETRY_SECONDS, attempt, EXCEL_MAX_RETRIES,
            )
            time.sleep(EXCEL_RETRY_SECONDS)
        except OSError as exc:
            raise StorageError(f"{what}失敗：{exc}") from exc
    raise StorageError(f"{what}失敗：Excel 檔一直被鎖住")


def _open_workbook(path: Path) -> Workbook:
    if path.exists():
        return load_workbook(path)
    workbook = Workbook()
    workbook.active.title = "紀錄"
    workbook.active.append(HEADERS)
    return workbook


def _as_datetime(value: object) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S"):
            try:
                return datetime.strptime(value, fmt)
            except ValueError:
                continue
    return None


def _as_kcal(value: object) -> int | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return round(value)
    return None


def _day_total(sheet, day: date) -> int:
    """該日所有已記錄餐點的熱量加總（整個頻道，這一版不分人）。每次重新加總，手動改過 Excel 也會跟著對。"""
    total = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        when = _as_datetime(row[0])
        kcal = _as_kcal(row[5]) if len(row) > 5 else None
        if when and when.date() == day and kcal:
            total += kcal
    return total


def format_ingredients(analysis: Analysis) -> str:
    return "、".join(f"{item.name} {item.weight_g}g {item.kcal}大卡" for item in analysis.ingredients)


def append_record(
    config: Config,
    when: datetime,
    uploader: str,
    analysis: Analysis,
    photo_path: Path,
) -> int:
    """在 Excel 最後面加一列，回傳含這一餐的「當日已攝取熱量」。永遠只加、不改、不刪既有列。"""
    relative_photo = os.path.relpath(photo_path, config.data_dir)

    def action() -> int:
        workbook = _open_workbook(config.excel_path)
        sheet = workbook.active
        previous = _day_total(sheet, when.date())
        meal_kcal = analysis.total_kcal if analysis.is_food else 0
        day_total = previous + meal_kcal

        row = sheet.max_row + 1
        sheet.cell(row=row, column=1, value=when).number_format = "yyyy/mm/dd hh:mm:ss"
        guards.write_text_cell(sheet.cell(row=row, column=2), uploader)
        if analysis.is_food:
            guards.write_text_cell(sheet.cell(row=row, column=3), analysis.name)
            guards.write_text_cell(sheet.cell(row=row, column=4), format_ingredients(analysis))
            sheet.cell(row=row, column=5, value=analysis.total_weight_g)
            sheet.cell(row=row, column=6, value=analysis.total_kcal)
        else:
            # 不是食物：食物敘述寫「不是食物」，成份欄寫判定出來的東西，重量與熱量留空
            guards.write_text_cell(sheet.cell(row=row, column=3), NOT_FOOD_LABEL)
            guards.write_text_cell(sheet.cell(row=row, column=4), analysis.name)
        sheet.cell(row=row, column=7, value=day_total)
        guards.write_text_cell(sheet.cell(row=row, column=8), relative_photo)

        # 先寫到暫存檔再替換正式檔，避免寫到一半當機把原檔毀掉
        temp_path = config.excel_path.with_suffix(".xlsx.tmp")
        workbook.save(temp_path)
        os.replace(temp_path, config.excel_path)
        return day_total

    return _with_retry(action, "寫入 Excel")


def meals_on(config: Config, day: date) -> list[Meal]:
    """某一天所有的餐點（含不是食物的），依時間排序。"""

    def action() -> list[Meal]:
        if not config.excel_path.exists():
            return []
        workbook = load_workbook(config.excel_path, read_only=True)
        try:
            sheet = workbook.active
            meals: list[Meal] = []
            for row in sheet.iter_rows(min_row=2, values_only=True):
                when = _as_datetime(row[0])
                if when and when.date() == day:
                    meals.append(Meal(when, str(row[1] or ""), str(row[2] or ""), _as_kcal(row[5])))
            return sorted(meals, key=lambda meal: meal.when)
        finally:
            workbook.close()

    return _with_retry(action, "讀取 Excel")
