from datetime import timedelta

import pytest
from openpyxl import load_workbook

from food_agent import storage
from food_agent.analyzer import Analysis, Ingredient

FOOD = Analysis(True, "雞腿便當", (Ingredient("白飯", 200, 280), Ingredient("雞腿", 150, 380)), 350, 660)
NOT_FOOD = Analysis(False, "馬克杯", (), 0, 0)


def rows(config):
    return list(load_workbook(config.excel_path).active.iter_rows(values_only=True))


def test_append_keeps_old_rows_and_accumulates(config, png_bytes, when):
    photo = storage.save_photo(config, png_bytes, when, "便當")
    assert storage.append_record(config, when, "小明", FOOD, photo) == 660
    assert storage.append_record(config, when + timedelta(hours=5), "小華", FOOD, photo) == 1320
    # 隔天重新計算
    assert storage.append_record(config, when + timedelta(days=1), "小明", FOOD, photo) == 660

    data = rows(config)
    assert data[0][0] == "日期時間" and len(data) == 4
    assert [row[6] for row in data[1:]] == [660, 1320, 660]
    assert data[1][3] == "白飯 200g 280大卡、雞腿 150g 380大卡"
    assert data[1][4:6] == (350, 660)


def test_existing_file_is_never_recreated(config, png_bytes, when):
    photo = storage.save_photo(config, png_bytes, when, "a")
    storage.append_record(config, when, "u", FOOD, photo)
    storage.append_record(config, when, "u", FOOD, photo)
    assert len(rows(config)) == 3


def test_not_food_row_and_no_calories(config, png_bytes, when):
    photo = storage.save_photo(config, png_bytes, when, storage.NOT_FOOD_LABEL)
    assert storage.append_record(config, when, "u", FOOD, photo) == 660
    assert storage.append_record(config, when, "u", NOT_FOOD, photo) == 660
    last = rows(config)[-1]
    assert last[2] == "不是食物" and last[3] == "馬克杯" and last[4] is None and last[5] is None and last[6] == 660


def test_formula_injection_stays_text(config, png_bytes, when):
    evil = Analysis(True, '=HYPERLINK("http://x","y")', (Ingredient("-cmd", 1, 1),), 1, 1)
    photo = storage.save_photo(config, png_bytes, when, "x")
    storage.append_record(config, when, "=1+1", evil, photo)
    sheet = load_workbook(config.excel_path).active
    for cell in (sheet["B2"], sheet["C2"], sheet["D2"]):
        assert cell.data_type != "f" and str(cell.value).startswith("'")


def test_photo_same_second_gets_suffix_and_never_overwrites(config, png_bytes, when):
    paths = [storage.save_photo(config, png_bytes, when, "牛排") for _ in range(3)]
    assert [p.name for p in paths] == [
        "2026-10-26 12-34-56 牛排.jpg",
        "2026-10-26 12-34-56 牛排-2.jpg",
        "2026-10-26 12-34-56 牛排-3.jpg",
    ]


def test_photo_name_cannot_escape_folder(config, png_bytes, when):
    path = storage.save_photo(config, png_bytes, when, "../../evil")
    assert path.parent == config.photos_dir
    assert not (config.data_dir.parent / "evil.jpg").exists()


def test_invalid_image_rejected():
    with pytest.raises(storage.StorageError):
        storage.validate_image(b"not an image")


def test_meals_on_filters_by_day(config, png_bytes, when):
    photo = storage.save_photo(config, png_bytes, when, "a")
    storage.append_record(config, when, "小明", FOOD, photo)
    storage.append_record(config, when + timedelta(days=1), "小明", FOOD, photo)
    assert len(storage.meals_on(config, when.date())) == 1
    assert storage.meals_on(config, (when + timedelta(days=5)).date()) == []


def test_locked_excel_retries_then_fails_without_touching_file(config, png_bytes, when, monkeypatch):
    monkeypatch.setattr(storage, "EXCEL_RETRY_SECONDS", 0)
    monkeypatch.setattr(storage, "EXCEL_MAX_RETRIES", 2)
    photo = storage.save_photo(config, png_bytes, when, "a")
    storage.append_record(config, when, "u", FOOD, photo)

    def locked(*args, **kwargs):
        raise PermissionError("locked")

    monkeypatch.setattr(storage.os, "replace", locked)
    with pytest.raises(storage.StorageError):
        storage.append_record(config, when, "u", FOOD, photo)
    assert len(rows(config)) == 2
